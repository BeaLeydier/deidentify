#!/usr/bin/env python3
"""scan_string_tails.py -- detect residual bytes hidden in string variables.

A fixed-width string field stores value + one \\0 terminator; the bytes AFTER the
terminator are NOT cleared by the tool and are invisible to normal reads (which
stop at the \\0), yet they are written to disk and travel with the file. They can
retain fragments of prior values -- i.e. real PII that a value-level scan misses.
This is a residual-data disclosure (MITRE CWE-212 / CWE-226).

This scanner reads a Stata .dta (format 117/118/119: Stata 13+) with NO external
tooling -- it parses the file's own <map> to locate the variable table and data,
then for every string variable inspects the bytes after each cell's terminator.

Usage:
    python3 scan_string_tails.py PATH.dta [OUT_DIR] [--xlsx=WORKBOOK.xlsx]

Outputs:
    string_tails_full.csv   one row per string cell that has residue (evidence)
    a `string_tails` sheet appended to the review workbook (per-variable summary)
Exit status is NONZERO if any residue is found (fail-loud), 0 if clean.
Requires openpyxl for the workbook sheet (pip install openpyxl).

Not Stata? The concept is identical for any fixed-width string storage; port the
"bytes after the terminator" check to your format. For pre-117 .dta, first re-save
as a modern format (Stata 13+: `save, replace`) -- residue is preserved by re-save.
"""
import csv, os, struct, sys
from collections import Counter

def parse_dta(raw):
    if raw[:11] != b"<stata_dta>":
        sys.exit("Not a format-117+ .dta (no <stata_dta>). Re-save as modern format first.")
    rel = int(raw[raw.index(b"<release>")+9: raw.index(b"</release>")])
    endian = "<" if raw[raw.index(b"<byteorder>")+11: raw.index(b"</byteorder>")] == b"LSF" else ">"
    kpos = raw.index(b"<K>") + 3
    nvar = struct.unpack_from(endian+"H", raw, kpos)[0]
    # <map>: 14 int64 absolute offsets; [2]=variable_types [3]=varnames [9]=data
    mpos = raw.index(b"<map>") + len(b"<map>")
    off = struct.unpack_from(endian + "14q", raw, mpos)
    # variable types: nvar uint16 codes after the tag
    tpos = off[2] + len(b"<variable_types>")
    codes = struct.unpack_from(endian + f"{nvar}H", raw, tpos)
    # variable names: fixed-width, 33 bytes (rel 117) or 129 bytes (rel 118/119)
    namelen = 33 if rel == 117 else 129
    npos = off[3] + len(b"<varnames>")
    names = [raw[npos+i*namelen: npos+(i+1)*namelen].split(b"\x00")[0].decode("latin1")
             for i in range(nvar)]
    def width(c):
        if 1 <= c <= 2045: return c            # str#  (code == byte length)
        return {32768:8, 65526:8, 65527:4, 65528:4, 65529:2, 65530:1}[c]  # strL,double,float,long,int,byte
    widths = [width(c) for c in codes]
    is_str = [1 <= c <= 2045 for c in codes]   # fixed-width strings only (strL handled elsewhere)
    data0 = off[9] + len(b"<data>")
    rec = sum(widths)
    nobs = (off[10] - data0) // rec if rec else 0   # off[10] = <strls> start
    return names, widths, is_str, data0, rec, nobs

def main():
    # usage: scan_string_tails.py PATH.dta [OUT_DIR] [--xlsx=WORKBOOK.xlsx]
    # Appends a `string_tails` sheet to the shared audit workbook (created if absent)
    # so the reviewer opens ONE file; the full per-cell dump stays a side CSV.
    pos = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not pos:
        sys.exit("usage: scan_string_tails.py PATH.dta [OUT_DIR] [--xlsx=WORKBOOK.xlsx]")
    path = pos[0]
    outdir = pos[1] if len(pos) > 1 else os.path.dirname(os.path.abspath(path)) or "."
    xlsx = None
    for a in sys.argv:
        if a.startswith("--xlsx="):
            xlsx = a.split("=", 1)[1]
    if xlsx is None:
        xlsx = f"{outdir}/audit_review.xlsx"
    raw = open(path, "rb").read()
    names, widths, is_str, data0, rec, nobs = parse_dta(raw)
    offs = [sum(widths[:i]) for i in range(len(widths))]

    full_rows, summary = [], []
    for vi, nm in enumerate(names):
        if not is_str[vi]:
            continue
        w = widths[vi]; o = offs[vi]
        n_res = 0; pats = Counter()
        for i in range(nobs):
            c = raw[data0 + i*rec + o: data0 + i*rec + o + w]
            z = c.find(b"\x00")
            if z == -1:            # full-width value, no terminator, no tail
                continue
            tail = c[z+1:]
            if tail == b"\x00" * len(tail):
                continue           # cleanly zero-padded
            n_res += 1
            pats[tail] += 1
            value = c[:z].decode("latin1")
            asc = "".join(chr(x) if 32 <= x < 127 else "." for x in tail)
            full_rows.append([nm, i+1, value, z, asc, tail.hex(" ")])
        example = ""
        if pats:
            t = pats.most_common(1)[0][0]
            example = "".join(chr(x) if 32 <= x < 127 else "." for x in t)
        summary.append([nm, f"str{w}", nobs, n_res,
                        round(100*n_res/nobs, 2) if nobs else 0, len(pats), example])

    # full per-cell evidence stays a side CSV (can be large); the review SUMMARY
    # goes into the shared audit workbook as a `string_tails` sheet.
    with open(f"{outdir}/string_tails_full.csv", "w", newline="") as f:
        wtr = csv.writer(f); wtr.writerow(["variable","obs","value","terminator_pos","tail_ascii","tail_hex"])
        wtr.writerows(full_rows)

    from openpyxl import Workbook, load_workbook
    if os.path.exists(xlsx):
        wb = load_workbook(xlsx)
    else:
        wb = Workbook(); wb.remove(wb.active)      # start clean; sheet added below
    if "string_tails" in wb.sheetnames:
        wb.remove(wb["string_tails"])
    ws = wb.create_sheet("string_tails")
    ws.append(["variable","type","n_cells","cells_with_residue","pct_residue",
               "distinct_patterns","example_residue"])
    for r in summary:
        ws.append(r)
    wb.save(xlsx)

    total = sum(r[3] for r in summary)
    print(f"scanned {sum(is_str)} string variable(s) x {nobs} obs in {os.path.basename(path)}")
    print(f"wrote string_tails_full.csv ({len(full_rows)} rows) and 'string_tails' sheet -> {xlsx}")
    if total:
        print(f"FLAG: {total} cell(s) carry hidden residual bytes -> strip-pii clean-overwrite required")
        sys.exit(1)
    print("OK: no string-tail residue found")

if __name__ == "__main__":
    main()
